import os
import json
import pandas as pd
from bs4 import BeautifulSoup # pyright: ignore[reportMissingImports]
from pypdf import PdfReader
from docx import Document # pyright: ignore[reportMissingImports]
from pptx import Presentation
from google import genai
import chromadb # pyright: ignore[reportMissingImports]
from dotenv import load_dotenv # pyright: ignore[reportMissingImports]

import sys

if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

load_dotenv()

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CHROMA_DATA_PATH = os.path.join(BASE_DIR, "chroma_db")
DATA_DIR = os.path.join(BASE_DIR, "data")

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY") or os.getenv("GOOGLE_API_KEY")
if not GEMINI_API_KEY:
    raise ValueError(
        "Chave de API não encontrada. Defina GEMINI_API_KEY ou GOOGLE_API_KEY no arquivo .env."
    )

genai_client = genai.Client(api_key=GEMINI_API_KEY)


def extract_text_from_file(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()
    text = ""
    try:
        if ext == ".pdf":
            reader = PdfReader(file_path)
            text = "\n".join([page.extract_text() or "" for page in reader.pages])
        elif ext == ".docx":
            try:
                doc = Document(file_path)
                lines = [p.text for p in doc.paragraphs if p.text]
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join(
                            [
                                cell.text.strip()
                                for cell in row.cells
                                if cell.text.strip()
                            ]
                        )
                        if row_text:
                            lines.append(row_text)
                text = "\n".join(lines)
            except Exception:
                import docx2txt # pyright: ignore[reportMissingImports]

                text = docx2txt.process(file_path) or ""
        elif ext in [".xlsx", ".xls"]:
            excel = pd.read_excel(file_path, sheet_name=None)
            for sheet, df in excel.items():
                text += f"\n--- Aba: {sheet} ---\n" + df.to_string(index=False)
        elif ext == ".csv":
            df = pd.read_csv(file_path)
            text = df.to_string(index=False)
        elif ext == ".pptx":
            prs = Presentation(file_path)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
        elif ext == ".json":
            with open(file_path, "r", encoding="utf-8") as f:
                text = json.dumps(json.load(f), indent=2, ensure_ascii=False)
        elif ext == ".html":
            with open(file_path, "r", encoding="utf-8") as f:
                soup = BeautifulSoup(f.read(), "html.parser")
                text = soup.get_text(separator="\n")
        elif ext in [".md", ".txt"]:
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.read()
    except Exception as e:
        print(f"Erro ao ler {file_path}: {e}")
    return text.strip()


def chunk_text(text: str, chunk_size: int = 700, overlap: int = 100) -> list[str]:
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        start += chunk_size - overlap
    return chunks


import time


def generate_embeddings_batch(
    texts: list[str], max_retries: int = 5
) -> list[list[float]]:
    for attempt in range(max_retries):
        try:
            response = genai_client.models.embed_content(
                model="gemini-embedding-001",
                contents=texts,
                config={"task_type": "RETRIEVAL_DOCUMENT"},
            )
            return [e.values for e in response.embeddings]
        except Exception as e:
            err_msg = str(e)
            if "429" in err_msg or "RESOURCE_EXHAUSTED" in err_msg:
                wait_time = (attempt + 1) * 3
                print(
                    f"Aviso: Taxa limite do Gemini atingida (429). Aguardando {wait_time}s antes de tentar novamente..."
                )
                time.sleep(wait_time)
            else:
                if attempt == max_retries - 1:
                    raise e
                time.sleep(2)
    return []


def generate_embedding(text: str) -> list[float]:
    res = generate_embeddings_batch([text])
    return res[0] if res else []


def run_ingestion(batch_size: int = 16):
    client = chromadb.PersistentClient(path=CHROMA_DATA_PATH)
    collection = client.get_or_create_collection(name="documentos_empresa")

    print(f"--- Iniciando ingestão otimizada de documentos ---")
    print(f"Diretório de documentos: {DATA_DIR}")
    print(f"Banco ChromaDB: {CHROMA_DATA_PATH}")
    total_docs = 0

    all_files = []
    for root, _, files in os.walk(DATA_DIR):
        for file in files:
            if file.startswith("."):
                continue
            all_files.append((root, file))

    print(f"Total de arquivos encontrados para processar: {len(all_files)}")

    for file_num, (root, file) in enumerate(all_files, 1):
        file_path = os.path.join(root, file)
        category = os.path.basename(root)
        print(f"\n[{file_num}/{len(all_files)}] Processando: {file} ({category})...")

        raw_text = extract_text_from_file(file_path)
        if not raw_text:
            print(f" -> Aviso: Nenhum texto extraído de {file}. Ignorando.")
            continue

        chunks = chunk_text(raw_text)
        print(
            f" -> Extraídos {len(chunks)} blocos de texto. Gerando embeddings em lotes..."
        )

        for b_start in range(0, len(chunks), batch_size):
            b_end = min(b_start + batch_size, len(chunks))
            batch_chunks = chunks[b_start:b_end]

            embs = generate_embeddings_batch(batch_chunks)
            if not embs:
                print(
                    f" -> Erro ao gerar embeddings para o lote {b_start}-{b_end}. Pulando..."
                )
                continue

            ids = [
                f"doc_{category}_{file}_{b_start + idx}"
                for idx in range(len(batch_chunks))
            ]
            metas = [
                {"source": file, "category": category, "chunk_index": b_start + idx}
                for idx in range(len(batch_chunks))
            ]

            collection.upsert(
                ids=ids, embeddings=embs, documents=batch_chunks, metadatas=metas
            )
            total_docs += len(batch_chunks)
            print(f"   [OK] Lote {b_start + 1}-{b_end} de {len(chunks)} indexado.")
            time.sleep(0.5)

    print("\n=======================================================")
    print(f"Sucesso: Ingestao concluida! Total de blocos indexados: {total_docs}")
    print("=======================================================")


if __name__ == "__main__":
    run_ingestion()
